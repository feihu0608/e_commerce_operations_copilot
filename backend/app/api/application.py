from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
import csv
import io
from pathlib import Path
import secrets
from typing import Any

from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel, EmailStr, Field
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from ..core.config import get_settings
from ..core.observability import RequestContextMiddleware, current_request_id
from ..core.security import create_access_token, get_current_user, hash_password, require_manager, verify_password
from ..domain.models import AuditLog, Competitor, ContentDocument, Experiment, GenerationTask, MetricRecord, PasswordReset, Product, TaskAttempt, TaskEvent, User, utcnow
from ..infrastructure.database import SessionLocal, get_db
from ..services.task_runtime import create_generation_task, record_audit, record_event


class LoginInput(BaseModel):
    account: str
    password: str


class RegisterInput(BaseModel):
    username: str = Field(min_length=3, max_length=60, pattern=r"^[A-Za-z0-9_\u4e00-\u9fa5]+$")
    email: EmailStr
    password: str = Field(min_length=8, max_length=128)


class ForgotInput(BaseModel):
    email: EmailStr


class ResetInput(BaseModel):
    token: str
    password: str = Field(min_length=8, max_length=128)


class ProductInput(BaseModel):
    name: str = Field(min_length=2, max_length=200)
    category: str = Field(min_length=2, max_length=100)
    price: Decimal = Field(gt=0)
    inventory: int = Field(ge=0)
    warning_threshold: int = Field(default=20, ge=0)
    summary: str = ""


class ContentUpdate(BaseModel):
    payload: dict[str, Any]


class TaskInput(BaseModel):
    kind: str = Field(pattern="^(image|video|diagnosis|creative|strategy|review)$")
    title: str = Field(min_length=2, max_length=200)


class DecisionInput(BaseModel):
    decision: str = Field(pattern="^(approved|rejected)$")
    note: str = Field(default="", max_length=500)


class ExperimentInput(BaseModel):
    product_id: int
    title: str = Field(min_length=2, max_length=200)
    audience: str = Field(min_length=2, max_length=300)
    channel: str = Field(min_length=2, max_length=50)
    budget: Decimal = Field(gt=0, le=1000000)
    period: str = Field(min_length=2, max_length=50)
    creative_angle: str = Field(min_length=2, max_length=300)
    target_ctr: Decimal = Field(gt=0, le=100)
    stop_roas: Decimal = Field(ge=0, le=100)


class MetricInput(BaseModel):
    period: str = Field(min_length=2, max_length=50)
    impressions: int = Field(ge=0)
    clicks: int = Field(ge=0)
    paid_orders: int = Field(ge=0)
    gmv: Decimal = Field(ge=0)
    ad_spend: Decimal = Field(ge=0)


def serialize_user(user: User):
    return {"id": user.id, "username": user.username, "email": user.email, "role": user.role}


def serialize_product(product: Product):
    image_url = product.image_url
    if not image_url or image_url == "/images/demo-phone.png":
        tile_names = [
            "曜石 X1", "Buds Air", "65W", "HUAWEI", "Apple", "Xiaomi", "OPPO",
            "vivo", "HONOR", "Samsung", "OnePlus", "DJI", "Sony",
        ]
        tile = next((index for index, name in enumerate(tile_names) if name.lower() in product.name.lower()), product.id % 13)
        image_url = f"/images/product-catalog-grid.png#tile={tile}"
    return {
        "id": product.id, "name": product.name, "category": product.category,
        "price": float(product.price), "status": product.status, "image_url": image_url,
        "summary": product.summary, "inventory": product.inventory,
        "warning_threshold": product.warning_threshold,
        "inventory_warning": product.inventory <= product.warning_threshold,
    }


def task_error_view(task: GenerationTask) -> tuple[str | None, bool]:
    message = task.error_message
    if not message:
        return None, True
    if "402 Payment Required" in message or "余额或额度不足" in message:
        return "硅基流动账户余额或额度不足，请充值后重新生成", False
    if "密钥无效" in message or "没有当前模型权限" in message:
        return "模型服务密钥无效或没有当前模型权限，请联系管理员检查配置", False
    return message, True


def serialize_task(task: GenerationTask):
    error_message, retryable = task_error_view(task)
    return {
        "id": task.id, "product_id": task.product_id, "kind": task.kind, "title": task.title,
        "status": task.status, "progress": task.progress, "provider_mode": task.provider_mode,
        "error_message": error_message, "retryable": retryable, "result_url": task.result_url,
        "request_id": task.request_id, "version": task.version, "retry_of_task_id": task.retry_of_task_id,
        "created_at": task.created_at, "updated_at": task.updated_at,
    }


def serialize_experiment(item: Experiment):
    return {
        "id": item.id, "product_id": item.product_id, "task_id": item.task_id, "title": item.title, "status": item.status,
        "strategy": item.strategy, "decision_note": item.decision_note,
        "created_by_id": item.created_by_id, "submitted_at": item.submitted_at,
        "reviewer_id": item.reviewer_id, "reviewed_at": item.reviewed_at,
        "created_at": item.created_at, "updated_at": item.updated_at,
    }


def metric_view(metric: MetricRecord):
    ctr = metric.clicks / metric.impressions if metric.impressions else None
    conversion = metric.paid_orders / metric.clicks if metric.clicks else None
    roas = float(metric.gmv / metric.ad_spend) if metric.ad_spend else None
    return {
        "id": metric.id, "period": metric.period, "impressions": metric.impressions,
        "clicks": metric.clicks, "paid_orders": metric.paid_orders, "gmv": float(metric.gmv),
        "ad_spend": float(metric.ad_spend), "ctr": None if ctr is None else round(ctr * 100, 2),
        "conversion_rate": None if conversion is None else round(conversion * 100, 2),
        "roas": None if roas is None else round(roas, 2),
    }


def seed_users(db: Session):
    settings = get_settings()
    if db.scalar(select(User).where(User.username == "operator")) is None:
        if not settings.demo_operator_password:
            raise RuntimeError("DEMO_OPERATOR_PASSWORD must be configured before first startup")
        db.add(User(username="operator", email="operator@example.com", password_hash=hash_password(settings.demo_operator_password), role="operator"))
    if db.scalar(select(User).where(User.username == "manager")) is None:
        if not settings.demo_manager_password:
            raise RuntimeError("DEMO_MANAGER_PASSWORD must be configured before first startup")
        db.add(User(username="manager", email="manager@example.com", password_hash=hash_password(settings.demo_manager_password), role="manager"))
    db.commit()


@asynccontextmanager
async def lifespan(_: FastAPI):
    Path(get_settings().storage_dir).mkdir(parents=True, exist_ok=True)
    with SessionLocal() as db:
        seed_users(db)
    yield


app = FastAPI(title="电商运营助手 API", version="1.0.0", lifespan=lifespan)
app.add_middleware(RequestContextMiddleware)
app.add_middleware(
    CORSMiddleware,
    allow_origins=get_settings().cors_origin_list,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.mount("/storage", StaticFiles(directory=get_settings().storage_dir, check_dir=False), name="storage")


@app.get("/api/health")
def health(db: Session = Depends(get_db)):
    db.execute(text("SELECT 1"))
    cfg = get_settings()
    return {"status": "ok", "database": "ok", "ai_mode": cfg.ai_mode, "media_mode": cfg.media_mode, "version": cfg.app_version}


@app.get("/api/ready")
def ready(db: Session = Depends(get_db)):
    from redis import Redis

    db.execute(text("SELECT 1"))
    client = Redis.from_url(get_settings().redis_url, socket_connect_timeout=2, socket_timeout=2)
    try:
        if not client.ping():
            raise RuntimeError("redis ping failed")
    finally:
        client.close()
    return {"status": "ready", "database": "ok", "redis": "ok"}


@app.post("/api/auth/login")
def login(body: LoginInput, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where((User.username == body.account) | (User.email == body.account)))
    if not user or not verify_password(body.password, user.password_hash):
        raise HTTPException(status_code=401, detail="账号或密码错误")
    return {"access_token": create_access_token(user), "token_type": "bearer", "user": serialize_user(user)}


@app.post("/api/auth/register", status_code=201)
def register(body: RegisterInput, db: Session = Depends(get_db)):
    if db.scalar(select(User).where((User.username == body.username) | (User.email == body.email))):
        raise HTTPException(status_code=409, detail="用户名或邮箱已存在")
    user = User(username=body.username, email=body.email, password_hash=hash_password(body.password), role="operator")
    db.add(user)
    db.commit()
    db.refresh(user)
    return {"access_token": create_access_token(user), "token_type": "bearer", "user": serialize_user(user)}


@app.post("/api/auth/forgot-password")
def forgot_password(body: ForgotInput, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.email == body.email))
    response: dict[str, Any] = {"message": "如果邮箱存在，重置请求已创建"}
    if user:
        raw = secrets.token_urlsafe(24)
        db.add(PasswordReset(user_id=user.id, token_hash=sha256(raw.encode()).hexdigest(), expires_at=datetime.now(timezone.utc) + timedelta(minutes=20)))
        db.commit()
        if get_settings().app_env == "development":
            response["debug_reset_token"] = raw
    return response


@app.post("/api/auth/reset-password")
def reset_password(body: ResetInput, db: Session = Depends(get_db)):
    item = db.scalar(select(PasswordReset).where(PasswordReset.token_hash == sha256(body.token.encode()).hexdigest()))
    now = datetime.now(timezone.utc)
    if not item or item.used_at or item.expires_at < now:
        raise HTTPException(status_code=400, detail="重置令牌无效或已过期")
    user = db.get(User, item.user_id)
    user.password_hash = hash_password(body.password)
    item.used_at = now
    db.commit()
    return {"message": "密码已更新"}


@app.get("/api/auth/me")
def me(user: User = Depends(get_current_user)):
    return serialize_user(user)


@app.post("/api/auth/logout")
def logout(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    user.token_version += 1
    record_audit(db, "auth.logout", "user", user.id, user.id)
    db.commit()
    return {"status": "logged_out"}


@app.get("/api/dashboard")
def dashboard(_: User = Depends(get_current_user), db: Session = Depends(get_db)):
    products = db.scalars(select(Product).order_by(Product.id)).all()
    metrics = db.scalars(select(MetricRecord)).all()
    warnings = [serialize_product(p) for p in products if p.inventory <= p.warning_threshold]
    gmv = sum((m.gmv for m in metrics), Decimal("0"))
    moving = len({m.product_id for m in metrics if m.paid_orders > 0})
    on_sale = sum(1 for p in products if p.status == "on_sale")
    pending = db.scalar(select(func.count()).select_from(Experiment).where(Experiment.status == "submitted")) or 0
    return {"product_count": len(products), "active_rate": round(moving / on_sale * 100, 1) if on_sale else 0,
            "gmv": float(gmv), "inventory_warnings": warnings, "pending_approvals": pending,
            "trend": [42, 51, 47, 63, 59, 72, 84]}


@app.get("/api/products")
def list_products(_: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return [serialize_product(p) for p in db.scalars(select(Product).order_by(Product.id)).all()]


@app.post("/api/products", status_code=201)
def create_product(body: ProductInput, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    product = Product(**body.model_dump(), image_url="/images/demo-phone.png")
    db.add(product)
    db.commit()
    db.refresh(product)
    return serialize_product(product)


@app.get("/api/products/{product_id}")
def product_detail(product_id: int, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(status_code=404, detail="商品不存在")
    competitors = db.scalars(select(Competitor).where(Competitor.product_id == product_id)).all()
    contents = db.scalars(select(ContentDocument).where(ContentDocument.product_id == product_id)).all()
    experiment = db.scalar(select(Experiment).where(Experiment.product_id == product_id).order_by(Experiment.id.desc()))
    metric = db.scalar(select(MetricRecord).where(MetricRecord.product_id == product_id).order_by(MetricRecord.id.desc()))
    return {
        "product": serialize_product(product),
        "competitors": [
            {
                "id": c.id,
                "name": c.name,
                "price": float(c.price),
                "highlights": c.highlights,
                "highlights_text": " · ".join(str(value) for value in (c.highlights or {}).values()),
            }
            for c in competitors
        ],
        "contents": [{"id": c.id, "type": c.content_type, "status": c.status, "revision": c.revision, "payload": c.payload} for c in contents],
        "experiment": None if not experiment else {"id": experiment.id, "title": experiment.title, "status": experiment.status, "strategy": experiment.strategy, "decision_note": experiment.decision_note},
        "metrics": None if not metric else metric_view(metric),
    }


def enqueue(
    db: Session,
    product_id: int,
    kind: str,
    title: str,
    actor_id: int | None = None,
    idempotency_key: str | None = None,
    retry_of_task_id: int | None = None,
) -> GenerationTask:
    cfg = get_settings()
    provider_mode = cfg.ai_mode if kind in ("diagnosis", "creative", "strategy", "review") else cfg.media_mode
    try:
        task, _ = create_generation_task(
            db, product_id, kind, title, provider_mode, actor_id, idempotency_key, retry_of_task_id
        )
        return task
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.post("/api/products/{product_id}/generate", status_code=202)
def generate(
    product_id: int,
    body: TaskInput,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    idempotency_key: str | None = Header(default=None, alias="Idempotency-Key"),
):
    return serialize_task(enqueue(db, product_id, body.kind, body.title, user.id, idempotency_key))


@app.get("/api/tasks")
def list_tasks(product_id: int | None = None, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    query = select(GenerationTask).order_by(GenerationTask.id.desc()).limit(50)
    if product_id:
        query = query.where(GenerationTask.product_id == product_id)
    return [serialize_task(t) for t in db.scalars(query).all()]


@app.get("/api/tasks/{task_id}/result")
def task_result(task_id: int, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    task = db.get(GenerationTask, task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    if task.status != "succeeded":
        raise HTTPException(status_code=409, detail="任务尚未生成成功")
    product = db.get(Product, task.product_id)
    content = None
    if task.kind in ("diagnosis", "creative", "review"):
        item = db.scalar(select(ContentDocument).where(ContentDocument.task_id == task.id))
        if item:
            content = {
                "id": item.id,
                "type": item.content_type,
                "status": item.status,
                "revision": item.revision,
                "payload": item.payload,
            }
    experiment = None
    if task.kind == "strategy":
        item = db.scalar(select(Experiment).where(Experiment.task_id == task.id))
        if item:
            experiment = serialize_experiment(item)
    return {
        "task": serialize_task(task),
        "product": None if not product else serialize_product(product),
        "content": content,
        "experiment": experiment,
        "result_url": task.result_url,
    }


@app.post("/api/tasks/{task_id}/cancel")
def cancel_task(task_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    task = db.get(GenerationTask, task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    if task.status in ("succeeded", "failed", "timeout", "cancelled"):
        raise HTTPException(status_code=409, detail="终态任务不能取消")
    task.status, task.progress = "cancelled", 100
    task.version += 1
    attempt = db.scalar(select(TaskAttempt).where(TaskAttempt.task_id == task.id, TaskAttempt.status == "running"))
    if attempt:
        attempt.status = "cancelled"
        attempt.error_code = "CANCEL_REQUESTED"
        attempt.finished_at = utcnow()
    record_event(db, task.id, "cancelled", actor_id=user.id)
    record_audit(db, "generation_task.cancel", "generation_task", task.id, user.id)
    db.commit()
    return serialize_task(task)


@app.post("/api/tasks/{task_id}/retry", status_code=202)
def retry_task(task_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    old = db.get(GenerationTask, task_id)
    if not old:
        raise HTTPException(status_code=404, detail="任务不存在")
    if old.status not in ("failed", "timeout", "cancelled"):
        raise HTTPException(status_code=409, detail="当前状态不能重试")
    _, retryable = task_error_view(old)
    if not retryable:
        raise HTTPException(status_code=409, detail="该失败需要先处理模型账户额度或权限，请处理后从原功能入口重新生成")
    return serialize_task(enqueue(db, old.product_id, old.kind, f"{old.title} 重试", user.id, retry_of_task_id=old.id))


@app.get("/api/tasks/{task_id}/events")
def task_events(task_id: int, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if not db.get(GenerationTask, task_id):
        raise HTTPException(status_code=404, detail="任务不存在")
    events = db.scalars(select(TaskEvent).where(TaskEvent.task_id == task_id).order_by(TaskEvent.id)).all()
    return [{"id": item.id, "type": item.event_type, "payload": item.payload, "created_at": item.created_at} for item in events]


@app.patch("/api/contents/{content_id}")
def update_content(content_id: int, body: ContentUpdate, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    item = db.get(ContentDocument, content_id)
    if not item:
        raise HTTPException(status_code=404, detail="内容不存在")
    item.payload, item.revision, item.status = body.payload, item.revision + 1, "draft"
    db.commit()
    return {"id": item.id, "revision": item.revision, "status": item.status, "payload": item.payload}


@app.post("/api/contents/{content_id}/confirm")
def confirm_content(content_id: int, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    item = db.get(ContentDocument, content_id)
    if not item:
        raise HTTPException(status_code=404, detail="内容不存在")
    item.status = "confirmed"
    db.commit()
    return {"id": item.id, "status": item.status, "revision": item.revision}


@app.get("/api/experiments")
def experiments(_: User = Depends(get_current_user), db: Session = Depends(get_db)):
    items = db.scalars(select(Experiment).order_by(Experiment.id.desc())).all()
    return [serialize_experiment(item) for item in items]


@app.post("/api/experiments", status_code=201)
def create_experiment(body: ExperimentInput, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if not db.get(Product, body.product_id):
        raise HTTPException(status_code=404, detail="商品不存在")
    item = Experiment(
        product_id=body.product_id,
        title=body.title,
        status="draft",
        created_by_id=getattr(user, "id", None),
        strategy={
            "audience": body.audience,
            "channel": body.channel,
            "budget": float(body.budget),
            "period": body.period,
            "creative_angle": body.creative_angle,
            "targets": {"ctr": float(body.target_ctr), "stop_roas": float(body.stop_roas)},
        },
    )
    db.add(item)
    db.flush()
    record_audit(db, "experiment.create", "experiment", item.id, getattr(user, "id", None))
    db.commit()
    db.refresh(item)
    return serialize_experiment(item)


@app.post("/api/experiments/{experiment_id}/submit")
def submit_experiment(experiment_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    item = db.get(Experiment, experiment_id)
    if not item:
        raise HTTPException(status_code=404, detail="投放方案不存在")
    if item.status != "draft":
        raise HTTPException(status_code=409, detail="只有草稿方案可以提交审批")
    item.status = "submitted"
    item.submitted_at = utcnow()
    record_audit(db, "experiment.submit", "experiment", item.id, getattr(user, "id", None))
    db.commit()
    return serialize_experiment(item)


@app.post("/api/experiments/{experiment_id}/decision")
def decide(experiment_id: int, body: DecisionInput, manager: User = Depends(require_manager), db: Session = Depends(get_db)):
    item = db.get(Experiment, experiment_id)
    if not item:
        raise HTTPException(status_code=404, detail="投放方案不存在")
    if item.status != "submitted":
        raise HTTPException(status_code=409, detail="方案已经处理")
    item.status = body.decision
    item.decision_note = body.note
    item.reviewer_id = manager.id
    item.reviewed_at = utcnow()
    record_audit(db, f"experiment.{body.decision}", "experiment", item.id, manager.id, note=body.note)
    db.commit()
    return {"id": item.id, "status": item.status, "decision_note": item.decision_note}


@app.post("/api/products/{product_id}/metrics", status_code=201)
def add_metric(product_id: int, body: MetricInput, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if not db.get(Product, product_id):
        raise HTTPException(status_code=404, detail="商品不存在")
    if body.clicks > body.impressions or body.paid_orders > body.clicks:
        raise HTTPException(status_code=422, detail="点击量不能超过曝光量，支付订单量不能超过点击量")
    item = MetricRecord(product_id=product_id, **body.model_dump())
    db.add(item)
    db.commit()
    db.refresh(item)
    return metric_view(item)


@app.get("/api/products/{product_id}/review")
def review(product_id: int, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    product = db.get(Product, product_id)
    metric = db.scalar(select(MetricRecord).where(MetricRecord.product_id == product_id).order_by(MetricRecord.id.desc()))
    experiment = db.scalar(select(Experiment).where(Experiment.product_id == product_id).order_by(Experiment.id.desc()))
    if not product or not metric:
        raise HTTPException(status_code=404, detail="商品或经营数据不足")
    values = metric_view(metric)
    findings = [
        f"CTR 为 {values['ctr']}%，影像卖点主图具备点击吸引力。" if values["ctr"] and values["ctr"] >= 3.2 else "CTR 未达到 3.2% 演示目标，需要调整首屏卖点。",
        f"转化率为 {values['conversion_rate']}%，需要结合详情页信任证据继续优化。",
    ]
    report = db.scalar(
        select(ContentDocument)
        .where(ContentDocument.product_id == product_id, ContentDocument.content_type == "review")
        .order_by(ContentDocument.id.desc())
    )
    return {
        "product": serialize_product(product),
        "experiment": None if not experiment else {"title": experiment.title, "status": experiment.status, "strategy": experiment.strategy},
        "metrics": values,
        "findings": findings,
        "actions": ["补充夜景人像原片与竞品对比", "突出 30 天无忧换机服务", "下一轮测试续航场景素材"],
        "ai_report": None if not report else {"id": report.id, "revision": report.revision, "status": report.status, "payload": report.payload},
    }


IMPORT_COLUMNS = [
    "统计日期", "商品编码", "品牌", "商品名称", "商品类目", "店铺", "渠道", "商品单价",
    "曝光量", "点击量", "支付订单量", "支付金额", "广告花费", "退款金额", "库存", "备注",
]
IMPORT_INTEGER_COLUMNS = ["曝光量", "点击量", "支付订单量", "库存"]
IMPORT_DECIMAL_COLUMNS = ["商品单价", "支付金额", "广告花费", "退款金额"]


def parse_import_file(file: UploadFile):
    name, raw = (file.filename or "").lower(), file.file.read()
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="演示版单文件不能超过 5 MB")
    if name.endswith(".csv"):
        try:
            reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
            columns = [str(x or "").strip() for x in (reader.fieldnames or [])]
            rows = list(reader)
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=422, detail="CSV 必须使用 UTF-8 编码") from exc
    elif name.endswith(".xlsx"):
        sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        values = list(sheet.iter_rows(values_only=True))
        columns = [str(x or "").strip() for x in values[0]] if values else []
        rows = [dict(zip(columns, row)) for row in values[1:]]
    else:
        raise HTTPException(status_code=422, detail="只支持 CSV 或 XLSX")
    rows = [row for row in rows if any(value not in (None, "") for value in row.values())]
    return columns, rows


def validate_import_rows(columns: list[str], rows: list[dict[str, Any]]):
    errors: list[str] = []
    missing_columns = [column for column in IMPORT_COLUMNS if column not in columns]
    if missing_columns:
        errors.append(f"缺少字段：{'、'.join(missing_columns)}")
    if not rows:
        errors.append("文件中没有可导入的数据行")
    normalized: list[dict[str, Any]] = []
    for index, source in enumerate(rows, start=2):
        if missing_columns:
            break
        row = {key: source.get(key) for key in IMPORT_COLUMNS}
        row_errors: list[str] = []
        for key in ["统计日期", "商品编码", "品牌", "商品名称", "商品类目", "店铺", "渠道"]:
            row[key] = str(row[key] or "").strip()
            if not row[key]:
                row_errors.append(f"{key}不能为空")
        raw_date = source.get("统计日期")
        try:
            if isinstance(raw_date, datetime):
                row["统计日期"] = raw_date.date().isoformat()
            elif isinstance(raw_date, date):
                row["统计日期"] = raw_date.isoformat()
            else:
                row["统计日期"] = datetime.strptime(str(raw_date).strip(), "%Y-%m-%d").date().isoformat()
        except (TypeError, ValueError):
            row_errors.append("统计日期必须为 YYYY-MM-DD")
        for key in IMPORT_INTEGER_COLUMNS:
            try:
                value = Decimal(str(source.get(key)).strip())
                if not value.is_finite() or value != value.to_integral_value() or value < 0:
                    raise ValueError
                row[key] = int(value)
            except (InvalidOperation, TypeError, ValueError):
                row_errors.append(f"{key}必须为非负整数")
        for key in IMPORT_DECIMAL_COLUMNS:
            try:
                value = Decimal(str(source.get(key)).strip())
                if not value.is_finite() or value < 0 or (key == "商品单价" and value <= 0):
                    raise ValueError
                row[key] = value
            except (InvalidOperation, TypeError, ValueError):
                row_errors.append(f"{key}必须为{'大于 0' if key == '商品单价' else '非负'}数值")
        if not row_errors and row["点击量"] > row["曝光量"]:
            row_errors.append("点击量不能超过曝光量")
        if not row_errors and row["支付订单量"] > row["点击量"]:
            row_errors.append("支付订单量不能超过点击量")
        if row_errors:
            errors.append(f"第 {index} 行：{'；'.join(row_errors)}")
        else:
            row["备注"] = str(source.get("备注") or "").strip()
            normalized.append(row)
    return normalized, errors


@app.post("/api/imports/preview")
def import_preview(file: UploadFile = File(...), _: User = Depends(get_current_user)):
    columns, rows = parse_import_file(file)
    _, errors = validate_import_rows(columns, rows)
    return {
        "status": "preview",
        "can_import": not errors,
        "message": "校验通过，尚未写入数据库" if not errors else "校验未通过，请修正文件后重试",
        "filename": file.filename,
        "row_count_total": len(rows),
        "row_count_previewed": min(len(rows), 20),
        "columns": columns,
        "rows": rows[:20],
        "validation_errors": errors,
    }


@app.post("/api/imports/commit")
def import_commit(file: UploadFile = File(...), _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    columns, rows = parse_import_file(file)
    normalized, errors = validate_import_rows(columns, rows)
    if errors:
        raise HTTPException(status_code=422, detail="；".join(errors[:10]))
    stats = {"products_created": 0, "products_updated": 0, "metrics_created": 0, "metrics_updated": 0}
    created_product_names: set[str] = set()
    updated_product_names: set[str] = set()
    try:
        for row in normalized:
            product = db.scalar(select(Product).where(Product.name == row["商品名称"]).order_by(Product.id))
            summary = f"{row['品牌']} · {row['商品编码']} · {row['店铺']}"
            if product:
                product.category = row["商品类目"]
                product.price = row["商品单价"]
                product.inventory = row["库存"]
                product.summary = summary
                if product.name not in created_product_names and product.name not in updated_product_names:
                    stats["products_updated"] += 1
                    updated_product_names.add(product.name)
            else:
                product = Product(
                    name=row["商品名称"], category=row["商品类目"], price=row["商品单价"],
                    inventory=row["库存"], summary=summary,
                )
                db.add(product)
                db.flush()
                stats["products_created"] += 1
                created_product_names.add(product.name)
            period = f"{row['统计日期']} · {row['渠道']}"
            metric = db.scalar(select(MetricRecord).where(MetricRecord.product_id == product.id, MetricRecord.period == period))
            if metric:
                metric.impressions = row["曝光量"]
                metric.clicks = row["点击量"]
                metric.paid_orders = row["支付订单量"]
                metric.gmv = row["支付金额"]
                metric.ad_spend = row["广告花费"]
                stats["metrics_updated"] += 1
            else:
                db.add(MetricRecord(
                    product_id=product.id, period=period, impressions=row["曝光量"], clicks=row["点击量"],
                    paid_orders=row["支付订单量"], gmv=row["支付金额"], ad_spend=row["广告花费"],
                ))
                stats["metrics_created"] += 1
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"status": "imported", "rows_imported": len(normalized), **stats}


@app.get("/api/settings")
def settings(_: User = Depends(require_manager)):
    cfg = get_settings()
    return {"ai_mode": cfg.ai_mode, "media_mode": cfg.media_mode, "base_url": cfg.siliconflow_base_url, "api_key_configured": bool(cfg.siliconflow_api_key), "models": {"text": cfg.text_model, "analysis": cfg.analysis_model, "image": cfg.image_model, "video_i2v": cfg.video_i2v_model, "video_t2v": cfg.video_t2v_model, "vision": cfg.vision_model}, "worker_concurrency": 1, "app_version": cfg.app_version}


@app.get("/api/audit-logs")
def audit_logs(_: User = Depends(require_manager), db: Session = Depends(get_db), limit: int = 100):
    items = db.scalars(select(AuditLog).order_by(AuditLog.id.desc()).limit(min(max(limit, 1), 500))).all()
    return [{"id": item.id, "actor_id": item.actor_id, "action": item.action, "resource_type": item.resource_type, "resource_id": item.resource_id, "request_id": item.request_id, "detail": item.detail, "created_at": item.created_at} for item in items]
