from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from hashlib import sha256
import csv
import io
import secrets
from typing import Any

from fastapi import Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from pydantic import BaseModel, EmailStr, Field
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from .config import get_settings
from .database import Base, SessionLocal, engine, get_db
from .models import Competitor, ContentDocument, Experiment, GenerationTask, MetricRecord, PasswordReset, Product, User, utcnow
from .security import create_access_token, get_current_user, hash_password, require_manager, verify_password
from .tasks import generate_content, generate_media


class LoginInput(BaseModel):
    account: str
    password: str


class RegisterInput(BaseModel):
    username: str = Field(min_length=3, max_length=60, pattern=r"^[A-Za-z0-9_\u4e00-\u9fa5]+$")
    email: EmailStr
    password: str = Field(min_length=8, max_length=128)


class ForgotInput(BaseModel):
    email: EmailStr


class ResetInput(BaseModel):
    token: str
    password: str = Field(min_length=8, max_length=128)


class ProductInput(BaseModel):
    name: str = Field(min_length=2, max_length=200)
    category: str = Field(min_length=2, max_length=100)
    price: Decimal = Field(gt=0)
    inventory: int = Field(ge=0)
    warning_threshold: int = Field(default=20, ge=0)
    summary: str = ""


class ContentUpdate(BaseModel):
    payload: dict[str, Any]


class TaskInput(BaseModel):
    kind: str = Field(pattern="^(image|video|diagnosis|creative)$")
    title: str = Field(min_length=2, max_length=200)


class DecisionInput(BaseModel):
    decision: str = Field(pattern="^(approved|rejected)$")
    note: str = Field(default="", max_length=500)


class MetricInput(BaseModel):
    period: str = Field(min_length=2, max_length=50)
    impressions: int = Field(ge=0)
    clicks: int = Field(ge=0)
    paid_orders: int = Field(ge=0)
    gmv: Decimal = Field(ge=0)
    ad_spend: Decimal = Field(ge=0)


def serialize_user(user: User):
    return {"id": user.id, "username": user.username, "email": user.email, "role": user.role}


def serialize_product(product: Product):
    return {
        "id": product.id, "name": product.name, "category": product.category,
        "price": float(product.price), "status": product.status, "image_url": product.image_url,
        "summary": product.summary, "inventory": product.inventory,
        "warning_threshold": product.warning_threshold,
        "inventory_warning": product.inventory <= product.warning_threshold,
    }


def serialize_task(task: GenerationTask):
    return {
        "id": task.id, "product_id": task.product_id, "kind": task.kind, "title": task.title,
        "status": task.status, "progress": task.progress, "provider_mode": task.provider_mode,
        "error_message": task.error_message, "result_url": task.result_url,
        "created_at": task.created_at, "updated_at": task.updated_at,
    }


def metric_view(metric: MetricRecord):
    ctr = metric.clicks / metric.impressions if metric.impressions else None
    conversion = metric.paid_orders / metric.clicks if metric.clicks else None
    roas = float(metric.gmv / metric.ad_spend) if metric.ad_spend else None
    return {
        "id": metric.id, "period": metric.period, "impressions": metric.impressions,
        "clicks": metric.clicks, "paid_orders": metric.paid_orders, "gmv": float(metric.gmv),
        "ad_spend": float(metric.ad_spend), "ctr": None if ctr is None else round(ctr * 100, 2),
        "conversion_rate": None if conversion is None else round(conversion * 100, 2),
        "roas": None if roas is None else round(roas, 2),
    }


def seed_users(db: Session):
    settings = get_settings()
    if db.scalar(select(User).where(User.username == "operator")) is None:
        if not settings.demo_operator_password:
            raise RuntimeError("DEMO_OPERATOR_PASSWORD must be configured before first startup")
        db.add(User(username="operator", email="operator@example.com", password_hash=hash_password(settings.demo_operator_password), role="operator"))
    if db.scalar(select(User).where(User.username == "manager")) is None:
        if not settings.demo_manager_password:
            raise RuntimeError("DEMO_MANAGER_PASSWORD must be configured before first startup")
        db.add(User(username="manager", email="manager@example.com", password_hash=hash_password(settings.demo_manager_password), role="manager"))
    db.commit()


@asynccontextmanager
async def lifespan(_: FastAPI):
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        seed_users(db)
    yield


app = FastAPI(title="电商运营助手 API", version="1.0.0", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/api/health")
def health(db: Session = Depends(get_db)):
    db.execute(text("SELECT 1"))
    cfg = get_settings()
    return {"status": "ok", "database": "ok", "ai_mode": cfg.ai_mode, "media_mode": cfg.media_mode}


@app.post("/api/auth/login")
def login(body: LoginInput, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where((User.username == body.account) | (User.email == body.account)))
    if not user or not verify_password(body.password, user.password_hash):
        raise HTTPException(status_code=401, detail="账号或密码错误")
    return {"access_token": create_access_token(user), "token_type": "bearer", "user": serialize_user(user)}


@app.post("/api/auth/register", status_code=201)
def register(body: RegisterInput, db: Session = Depends(get_db)):
    if db.scalar(select(User).where((User.username == body.username) | (User.email == body.email))):
        raise HTTPException(status_code=409, detail="用户名或邮箱已存在")
    user = User(username=body.username, email=body.email, password_hash=hash_password(body.password), role="operator")
    db.add(user)
    db.commit()
    db.refresh(user)
    return {"access_token": create_access_token(user), "token_type": "bearer", "user": serialize_user(user)}


@app.post("/api/auth/forgot-password")
def forgot_password(body: ForgotInput, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.email == body.email))
    response: dict[str, Any] = {"message": "如果邮箱存在，重置请求已创建"}
    if user:
        raw = secrets.token_urlsafe(24)
        db.add(PasswordReset(user_id=user.id, token_hash=sha256(raw.encode()).hexdigest(), expires_at=datetime.now(timezone.utc) + timedelta(minutes=20)))
        db.commit()
        if get_settings().app_env == "development":
            response["debug_reset_token"] = raw
    return response


@app.post("/api/auth/reset-password")
def reset_password(body: ResetInput, db: Session = Depends(get_db)):
    item = db.scalar(select(PasswordReset).where(PasswordReset.token_hash == sha256(body.token.encode()).hexdigest()))
    now = datetime.now(timezone.utc)
    if not item or item.used_at or item.expires_at < now:
        raise HTTPException(status_code=400, detail="重置令牌无效或已过期")
    user = db.get(User, item.user_id)
    user.password_hash = hash_password(body.password)
    item.used_at = now
    db.commit()
    return {"message": "密码已更新"}


@app.get("/api/auth/me")
def me(user: User = Depends(get_current_user)):
    return serialize_user(user)


@app.get("/api/dashboard")
def dashboard(_: User = Depends(get_current_user), db: Session = Depends(get_db)):
    products = db.scalars(select(Product).order_by(Product.id)).all()
    metrics = db.scalars(select(MetricRecord)).all()
    warnings = [serialize_product(p) for p in products if p.inventory <= p.warning_threshold]
    gmv = sum((m.gmv for m in metrics), Decimal("0"))
    moving = len({m.product_id for m in metrics if m.paid_orders > 0})
    on_sale = sum(1 for p in products if p.status == "on_sale")
    pending = db.scalar(select(func.count()).select_from(Experiment).where(Experiment.status == "submitted")) or 0
    return {"product_count": len(products), "active_rate": round(moving / on_sale * 100, 1) if on_sale else 0,
            "gmv": float(gmv), "inventory_warnings": warnings, "pending_approvals": pending,
            "trend": [42, 51, 47, 63, 59, 72, 84]}


@app.get("/api/products")
def list_products(_: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return [serialize_product(p) for p in db.scalars(select(Product).order_by(Product.id)).all()]


@app.post("/api/products", status_code=201)
def create_product(body: ProductInput, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    product = Product(**body.model_dump(), image_url="/images/demo-phone.png")
    db.add(product)
    db.commit()
    db.refresh(product)
    return serialize_product(product)


@app.get("/api/products/{product_id}")
def product_detail(product_id: int, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(status_code=404, detail="商品不存在")
    competitors = db.scalars(select(Competitor).where(Competitor.product_id == product_id)).all()
    contents = db.scalars(select(ContentDocument).where(ContentDocument.product_id == product_id)).all()
    experiment = db.scalar(select(Experiment).where(Experiment.product_id == product_id).order_by(Experiment.id.desc()))
    metric = db.scalar(select(MetricRecord).where(MetricRecord.product_id == product_id).order_by(MetricRecord.id.desc()))
    return {
        "product": serialize_product(product),
        "competitors": [
            {
                "id": c.id,
                "name": c.name,
                "price": float(c.price),
                "highlights": c.highlights,
                "highlights_text": " · ".join(str(value) for value in (c.highlights or {}).values()),
            }
            for c in competitors
        ],
        "contents": [{"id": c.id, "type": c.content_type, "status": c.status, "revision": c.revision, "payload": c.payload} for c in contents],
        "experiment": None if not experiment else {"id": experiment.id, "title": experiment.title, "status": experiment.status, "strategy": experiment.strategy, "decision_note": experiment.decision_note},
        "metrics": None if not metric else metric_view(metric),
    }


def enqueue(db: Session, product_id: int, kind: str, title: str) -> GenerationTask:
    if not db.get(Product, product_id):
        raise HTTPException(status_code=404, detail="商品不存在")
    cfg = get_settings()
    provider_mode = cfg.ai_mode if kind in ("diagnosis", "creative") else cfg.media_mode
    task = GenerationTask(product_id=product_id, kind=kind, title=title, provider_mode=provider_mode)
    db.add(task)
    db.commit()
    db.refresh(task)
    (generate_content.delay(task.id, kind) if kind in ("diagnosis", "creative") else generate_media.delay(task.id))
    return task


@app.post("/api/products/{product_id}/generate", status_code=202)
def generate(product_id: int, body: TaskInput, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return serialize_task(enqueue(db, product_id, body.kind, body.title))


@app.get("/api/tasks")
def list_tasks(product_id: int | None = None, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    query = select(GenerationTask).order_by(GenerationTask.id.desc()).limit(50)
    if product_id:
        query = query.where(GenerationTask.product_id == product_id)
    return [serialize_task(t) for t in db.scalars(query).all()]


@app.post("/api/tasks/{task_id}/cancel")
def cancel_task(task_id: int, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    task = db.get(GenerationTask, task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    if task.status in ("succeeded", "failed", "timeout", "cancelled"):
        raise HTTPException(status_code=409, detail="终态任务不能取消")
    task.status, task.progress = "cancelled", 100
    db.commit()
    return serialize_task(task)


@app.post("/api/tasks/{task_id}/retry", status_code=202)
def retry_task(task_id: int, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    old = db.get(GenerationTask, task_id)
    if not old:
        raise HTTPException(status_code=404, detail="任务不存在")
    if old.status not in ("failed", "timeout", "cancelled"):
        raise HTTPException(status_code=409, detail="当前状态不能重试")
    return serialize_task(enqueue(db, old.product_id, old.kind, f"{old.title} 重试"))


@app.patch("/api/contents/{content_id}")
def update_content(content_id: int, body: ContentUpdate, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    item = db.get(ContentDocument, content_id)
    if not item:
        raise HTTPException(status_code=404, detail="内容不存在")
    item.payload, item.revision, item.status = body.payload, item.revision + 1, "draft"
    db.commit()
    return {"id": item.id, "revision": item.revision, "status": item.status, "payload": item.payload}


@app.post("/api/contents/{content_id}/confirm")
def confirm_content(content_id: int, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    item = db.get(ContentDocument, content_id)
    if not item:
        raise HTTPException(status_code=404, detail="内容不存在")
    item.status = "confirmed"
    db.commit()
    return {"id": item.id, "status": item.status, "revision": item.revision}


@app.get("/api/experiments")
def experiments(_: User = Depends(get_current_user), db: Session = Depends(get_db)):
    items = db.scalars(select(Experiment).order_by(Experiment.id.desc())).all()
    return [{"id": x.id, "product_id": x.product_id, "title": x.title, "status": x.status, "strategy": x.strategy, "decision_note": x.decision_note} for x in items]


@app.post("/api/experiments/{experiment_id}/decision")
def decide(experiment_id: int, body: DecisionInput, manager: User = Depends(require_manager), db: Session = Depends(get_db)):
    item = db.get(Experiment, experiment_id)
    if not item:
        raise HTTPException(status_code=404, detail="投放方案不存在")
    if item.status != "submitted":
        raise HTTPException(status_code=409, detail="方案已经处理")
    item.status = body.decision
    item.decision_note = body.note
    item.reviewer_id = manager.id
    item.reviewed_at = utcnow()
    db.commit()
    return {"id": item.id, "status": item.status, "decision_note": item.decision_note}


@app.post("/api/products/{product_id}/metrics", status_code=201)
def add_metric(product_id: int, body: MetricInput, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if not db.get(Product, product_id):
        raise HTTPException(status_code=404, detail="商品不存在")
    if body.clicks > body.impressions or body.paid_orders > body.clicks:
        raise HTTPException(status_code=422, detail="点击量不能超过曝光量，支付订单量不能超过点击量")
    item = MetricRecord(product_id=product_id, **body.model_dump())
    db.add(item)
    db.commit()
    db.refresh(item)
    return metric_view(item)


@app.get("/api/products/{product_id}/review")
def review(product_id: int, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    product = db.get(Product, product_id)
    metric = db.scalar(select(MetricRecord).where(MetricRecord.product_id == product_id).order_by(MetricRecord.id.desc()))
    experiment = db.scalar(select(Experiment).where(Experiment.product_id == product_id).order_by(Experiment.id.desc()))
    if not product or not metric:
        raise HTTPException(status_code=404, detail="商品或经营数据不足")
    values = metric_view(metric)
    findings = [
        f"CTR 为 {values['ctr']}%，影像卖点主图具备点击吸引力。" if values["ctr"] and values["ctr"] >= 3.2 else "CTR 未达到 3.2% 演示目标，需要调整首屏卖点。",
        f"转化率为 {values['conversion_rate']}%，需要结合详情页信任证据继续优化。",
    ]
    return {"product": serialize_product(product), "experiment": None if not experiment else {"title": experiment.title, "status": experiment.status, "strategy": experiment.strategy}, "metrics": values, "findings": findings, "actions": ["补充夜景人像原片与竞品对比", "突出 30 天无忧换机服务", "下一轮测试续航场景素材"]}


@app.post("/api/imports/preview")
def import_preview(file: UploadFile = File(...), _: User = Depends(get_current_user)):
    name, raw = (file.filename or "").lower(), file.file.read()
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="演示版单文件不能超过 5 MB")
    if name.endswith(".csv"):
        rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))[:20]
    elif name.endswith(".xlsx"):
        sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        values = list(sheet.iter_rows(values_only=True))
        headers = [str(x or "") for x in values[0]] if values else []
        rows = [dict(zip(headers, row)) for row in values[1:21]]
    else:
        raise HTTPException(status_code=422, detail="只支持 CSV 或 XLSX")
    return {"filename": file.filename, "row_count_previewed": len(rows), "columns": list(rows[0].keys()) if rows else [], "rows": rows, "errors": []}


@app.get("/api/settings")
def settings(_: User = Depends(require_manager)):
    cfg = get_settings()
    return {"ai_mode": cfg.ai_mode, "media_mode": cfg.media_mode, "base_url": cfg.siliconflow_base_url, "api_key_configured": bool(cfg.siliconflow_api_key), "models": {"text": cfg.text_model, "analysis": cfg.analysis_model, "image": cfg.image_model, "video_i2v": cfg.video_i2v_model, "video_t2v": cfg.video_t2v_model, "vision": cfg.vision_model}, "worker_concurrency": 1}
